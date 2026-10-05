"""Combine the per-state workbooks in data/ into one all-India workbook.

Run after build_villages_excel.py:  python scripts/build_all_india_excel.py
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DATA = Path(__file__).resolve().parent.parent / "data"
OUT = DATA / "All_India_Villages_Pincodes.xlsx"
SUB_TYPES = ["Mandal", "Tehsil", "Taluk", "Taluka", "Block", "Circle", "Subdivision"]


def normalise(df):
    """Rename the state-specific Mandal/Tehsil/... columns to a common Sub-District column."""
    sub = next(t for t in SUB_TYPES if t in df.columns)
    df = df.rename(columns={sub: "Sub-District", f"{sub} LGD Code": "Sub-District LGD Code"})
    df.insert(df.columns.get_loc("Sub-District"), "Sub-District Type", sub)
    return df.drop(columns="S.No.")


def style(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 2000)), 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)


villages, subs, summary, pins, district_counts = [], [], [], [], {}
for f in sorted(DATA.glob("*_Villages_Pincodes.xlsx")):
    if f == OUT:
        continue
    sheets = pd.read_excel(f, sheet_name=None, dtype=str)
    for name, df in sheets.items():
        if name == "Summary":
            # district totals come from LGD's district list (some districts have no sub-districts)
            for _, r in df.dropna(subset=["Districts"]).iterrows():
                district_counts[r["State"]] = int(r["Districts"])
            continue
        if name.endswith(" Villages"):
            villages.append(normalise(df))
        elif name.endswith(" PIN Codes"):
            pins.append(df.drop(columns="S.No."))
        else:
            subs.append(normalise(df))

v = pd.concat(villages, ignore_index=True)
v = v.sort_values(["State", "District", "Sub-District", "Village"], key=lambda c: c.str.lower()).reset_index(drop=True)
v.insert(0, "S.No.", range(1, len(v) + 1))
s = pd.concat(subs, ignore_index=True)
s["No. of Villages"] = s["No. of Villages"].astype(int)
s = s.sort_values(["State", "District", "Sub-District"], key=lambda c: c.str.lower() if c.dtype == object else c).reset_index(drop=True)
s.insert(0, "S.No.", range(1, len(s) + 1))

for state, g in s.groupby("State", sort=True):
    vg = v[v["State"] == state]
    summary.append({
        "State / UT": state,
        "Sub-District Type": g["Sub-District Type"].iloc[0],
        "Districts": district_counts[state],
        "Sub-Districts": len(g),
        "Villages": len(vg),
        "Villages with Pincode": int(vg["Pincode"].notna().sum()),
        "Distinct Pincodes": int(vg["Pincode"].nunique()),
    })
sm = pd.DataFrame(summary)
total = sm.drop(columns=["State / UT", "Sub-District Type"]).sum()
total["Distinct Pincodes"] = v["Pincode"].nunique()
sm = pd.concat([sm, pd.DataFrame([{"State / UT": "ALL INDIA", "Sub-District Type": "", **total}])], ignore_index=True)

for df in (v, s):
    for c in df.columns:
        if c.endswith("Code") or c == "Pincode":
            df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")

note = pd.DataFrame({"Source": [
    "Local Government Directory (lgdirectory.gov.in), Govt. of India - data as on 02Oct2026",
    "Pincodes: LGD village-to-pincode mapping (blank = not mapped in LGD)",
    "Sub-District = Mandal / Tehsil / Taluk / Taluka / Block / Circle / Subdivision, as used in each state",
]})
with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    sm.to_excel(xw, sheet_name="Summary", index=False)
    note.to_excel(xw, sheet_name="Summary", index=False, startrow=len(sm) + 3)
    v.to_excel(xw, sheet_name="All India Villages", index=False)
    s.to_excel(xw, sheet_name="Sub-Districts", index=False)
    if pins:
        p = pd.concat(pins, ignore_index=True)
        p.insert(0, "S.No.", range(1, len(p) + 1))
        p.to_excel(xw, sheet_name="Urban PIN Codes (No Villages)", index=False)
    for ws in xw.book.worksheets:
        style(ws)
    xw.book["Summary"].cell(row=len(sm) + 1, column=1).font = Font(bold=True)
print(sm.to_string(index=False))
