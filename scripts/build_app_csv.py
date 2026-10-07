"""Export District / Mandal / Village tables as linked CSVs in the app's format.

Matches the app's CSV layout: "state_id","state_name","country_id","country_name","id","<name>","<parent>_id".
Each level has its own running id (1, 2, 3 ...), assigned alphabetically within the parent so every
parent's children sit in one continuous block; each row carries its parent's id as the link.

Usage: python scripts/build_app_csv.py
"""
import csv
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

DATA = Path(__file__).resolve().parent.parent / "data"
COUNTRY_ID, COUNTRY = 1, "India"
# (state_id in the app, state name, source workbook, villages sheet, sub-districts sheet, sub label, file prefix)
STATES = [
    (2, "Andhra Pradesh", "AP_Telangana_Villages_Pincodes.xlsx", "AP Villages", "AP Mandals", "Mandal", "ap"),
    (32, "Telangana", "AP_Telangana_Villages_Pincodes.xlsx", "Telangana Villages", "Telangana Mandals", "Mandal", "ts"),
]
# ids keep running across states in the order above, so every district/mandal/village id is unique in India
next_id = {"district": 1, "sub": 1, "village": 1}


def number(df, sort_cols, start=1):
    df = df.sort_values(sort_cols, key=lambda c: c.str.lower() if c.dtype == object else c, kind="stable")
    return df.assign(id=range(start, start + len(df)))


def write(df, path):
    df.to_csv(path, index=False, quoting=csv.QUOTE_ALL)


for state_id, state, src, vsheet, ssheet, sub, prefix in STATES:
    v = pd.read_excel(DATA / src, sheet_name=vsheet, dtype=str)
    s = pd.read_excel(DATA / src, sheet_name=ssheet, dtype=str)
    base = {"state_id": state_id, "state_name": state, "country_id": COUNTRY_ID, "country_name": COUNTRY}
    sl = sub.lower()

    d = number(s[["District", "District LGD Code"]].drop_duplicates("District LGD Code"), ["District"],
               next_id["district"])
    dmap = d.set_index("District LGD Code")["id"]

    m = s.assign(district_id=s["District LGD Code"].map(dmap))
    m = number(m, ["district_id", sub], next_id["sub"])
    mmap = m.set_index(f"{sub} LGD Code")["id"]

    vv = v.assign(**{f"{sl}_id": v[f"{sub} LGD Code"].map(mmap)})
    vv = number(vv, [f"{sl}_id", "Village"], next_id["village"])
    next_id.update(district=d["id"].max() + 1, sub=m["id"].max() + 1, village=vv["id"].max() + 1)

    out = DATA / "csv" / prefix
    out.mkdir(parents=True, exist_ok=True)
    write(pd.DataFrame({**base, "id": d["id"], "district_name": d["District"]}), out / f"{prefix}_districts.csv")
    write(pd.DataFrame({**base, "id": m["id"], f"{sl}_name": m[sub], "district_id": m["district_id"]}),
          out / f"{prefix}_{sl}s.csv")
    write(pd.DataFrame({**base, "id": vv["id"], "village_name": vv["Village"], f"{sl}_id": vv[f"{sl}_id"],
                        "pincode": vv["Pincode"].fillna("")}), out / f"{prefix}_villages.csv")

    # one sheet with every level and its id side by side
    mi = m.set_index("id")
    di = d.set_index("id")
    dist_id = vv[f"{sl}_id"].map(mi["district_id"])
    flat = pd.DataFrame({
        "country_id": COUNTRY_ID, "country_name": COUNTRY,
        "state_id": state_id, "state_name": state,
        "district_id": dist_id, "district_name": dist_id.map(di["District"]),
        f"{sl}_id": vv[f"{sl}_id"], f"{sl}_name": vv[f"{sl}_id"].map(mi[sub]),
        "village_id": vv["id"], "village_name": vv["Village"],
        "pincode": vv["Pincode"].fillna(""),
    })
    write(flat, out / f"{prefix}_all_levels.csv")
    xl = flat.copy()
    for c in [c for c in xl.columns if c.endswith("_id")] + ["pincode"]:
        xl[c] = pd.to_numeric(xl[c], errors="coerce").astype("Int64")
    with pd.ExcelWriter(out / f"{prefix}_all_levels.xlsx", engine="openpyxl") as xw:
        xl.to_excel(xw, sheet_name=f"{state[:31]}", index=False)
        ws = xw.sheets[f"{state[:31]}"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col[:2000]) + 2))
    assert (flat["district_name"] == vv["District"].values).all() and (flat[f"{sl}_name"] == vv[sub].values).all()

    # linked workbook: one sheet per level, each row holds its parent's id
    linked = {
        "1 Country": pd.DataFrame({"country_id": [COUNTRY_ID], "country_name": [COUNTRY]}),
        "2 State": pd.DataFrame({"state_id": [state_id], "state_name": [state], "country_id": [COUNTRY_ID]}),
        "3 Districts": pd.DataFrame({"district_id": d["id"], "district_name": d["District"], "state_id": state_id}),
        f"4 {sub}s": pd.DataFrame({f"{sl}_id": m["id"], f"{sl}_name": m[sub], "district_id": m["district_id"]}),
        "5 Villages": pd.DataFrame({"village_id": vv["id"], "village_name": vv["Village"],
                                    f"{sl}_id": vv[f"{sl}_id"],
                                    "pincode": pd.to_numeric(vv["Pincode"], errors="coerce").astype("Int64")}),
    }
    linked_csv = {"1 Country": "1_country", "2 State": "2_state", "3 Districts": "3_districts",
                  f"4 {sub}s": f"4_{sl}s", "5 Villages": "5_villages"}
    for name, df in linked.items():
        write(df, out / f"{prefix}_linked_{linked_csv[name]}.csv")
    # the same five sheets in one CSV file: each sheet is a section (sheet name, header, rows)
    with open(out / f"{prefix}_linked.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, quoting=csv.QUOTE_ALL)
        for i, (name, df) in enumerate(linked.items()):
            if i:
                w.writerow([])
            w.writerow([f"Sheet: {name}"])
            ids = [c for c in df.columns if c.endswith("_id")]
            w.writerow([f"{c} (link)" if c in ids[1:] else c for c in df.columns])
            w.writerows(df.astype("string").fillna("").itertuples(index=False))
    # each Excel tab as its own CSV, same headers as the tabs, all zipped into one download
    tab_dir = out / f"{prefix}_linked_csv"
    tab_dir.mkdir(exist_ok=True)
    with zipfile.ZipFile(out / f"{prefix}_linked_csv.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for name, df in linked.items():
            ids = [c for c in df.columns if c.endswith("_id")]
            tab = df.rename(columns={c: f"{c} (link)" for c in ids[1:]})
            write(tab, tab_dir / f"{name}.csv")
            z.write(tab_dir / f"{name}.csv", f"{name}.csv")
    link_fill = PatternFill("solid", fgColor="FFF2CC")
    with pd.ExcelWriter(out / f"{prefix}_linked.xlsx", engine="openpyxl") as xw:
        for name, df in linked.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                head = col[0]
                head.font = Font(bold=True, color="FFFFFF")
                head.fill = PatternFill("solid", fgColor="1F4E78")
                ws.column_dimensions[head.column_letter].width = max(14, min(40, max(len(str(c.value or "")) for c in col[:2000]) + 2))
            # highlight the link (parent id) column: last id column after the name
            parent_col = [c for c in df.columns if c.endswith("_id")][1:]
            for c in parent_col:
                idx = list(df.columns).index(c) + 1
                for row in ws.iter_rows(min_row=1, min_col=idx, max_col=idx):
                    for cell in row:
                        if cell.row > 1:
                            cell.fill = link_fill
                ws.cell(row=1, column=idx).value = f"{c} (link)"

    # single sheet: every level's id + name, each followed by the link to its parent
    single = pd.DataFrame({
        "country_id": COUNTRY_ID, "country_name": COUNTRY,
        "state_id": state_id, "state_name": state, "state.country_id (link)": COUNTRY_ID,
        "district_id": flat["district_id"], "district_name": flat["district_name"],
        "district.state_id (link)": state_id,
        f"{sl}_id": flat[f"{sl}_id"], f"{sl}_name": flat[f"{sl}_name"],
        f"{sl}.district_id (link)": flat["district_id"],
        "village_id": flat["village_id"], "village_name": flat["village_name"],
        f"village.{sl}_id (link)": flat[f"{sl}_id"],
        "pincode": flat["pincode"],
    })
    write(single, out / f"{prefix}_single_sheet.csv")
    xs = single.copy()
    for c in xs.columns:
        if "_id" in c or c == "pincode":
            xs[c] = pd.to_numeric(xs[c], errors="coerce").astype("Int64")
    with pd.ExcelWriter(out / f"{prefix}_single_sheet.xlsx", engine="openpyxl") as xw:
        xs.to_excel(xw, sheet_name=state[:31], index=False)
        ws = xw.sheets[state[:31]]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(xs.columns, 1):
            head = ws.cell(row=1, column=i)
            head.font = Font(bold=True, color="FFFFFF")
            head.fill = PatternFill("solid", fgColor="C55A11" if "(link)" in c else "1F4E78")
            ws.column_dimensions[head.column_letter].width = max(12, min(36, len(c) + 2,
                max(len(str(v)) for v in xs[c].head(2000).astype(str)) + 4))
        for i, c in enumerate(xs.columns, 1):
            if "(link)" in c:
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].fill = link_fill

    # checks: every link points at an existing parent, and names line up with the source
    assert vv[f"{sl}_id"].notna().all() and m["district_id"].notna().all()
    chk = vv.merge(m[["id", sub, "district_id"]].rename(columns={"id": f"{sl}_id", sub: "_m"}), on=f"{sl}_id")
    chk = chk.merge(d[["id", "District"]].rename(columns={"id": "district_id", "District": "_d"}), on="district_id")
    assert (chk["_m"] == chk[sub]).all() and (chk["_d"] == chk["District"]).all()
    print(prefix, "districts", len(d), f"{sl}s", len(m), "villages", len(vv))
