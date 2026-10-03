"""Build State > District > Mandal/Tehsil > Village Excel sheets with PIN codes.

Source: Local Government Directory (lgdirectory.gov.in), via the daily LGD
archive mirror at https://github.com/ramSeraph/opendata (release "lgd-latest").

Usage: python scripts/build_villages_excel.py <dir with extracted LGD csvs> <date e.g. 02Oct2026> [workbook.xlsx ...]
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC, DATE, ONLY = Path(sys.argv[1]), sys.argv[2], sys.argv[3:]
OUT = Path(__file__).resolve().parent.parent / "data"

# Only the workbooks named on the command line are rebuilt (all if none given).
# (output file, [(state code, sheet prefix, sub-district label)])
WORKBOOKS = [
    ("AP_Telangana_Villages_Pincodes.xlsx", [("28", "AP", "Mandal"), ("36", "Telangana", "Mandal")]),
    ("Odisha_Villages_Pincodes.xlsx", [("21", "Odisha", "Tehsil")]),
    ("TamilNadu_Villages_Pincodes.xlsx", [("33", "Tamil Nadu", "Taluk")]),
    ("Karnataka_Villages_Pincodes.xlsx", [("29", "Karnataka", "Taluk")]),
    ("Kerala_Villages_Pincodes.xlsx", [("32", "Kerala", "Taluk")]),
    ("Maharashtra_Villages_Pincodes.xlsx", [("27", "Maharashtra", "Taluka")]),
    ("Gujarat_Villages_Pincodes.xlsx", [("24", "Gujarat", "Taluka")]),
    ("Rajasthan_Villages_Pincodes.xlsx", [("8", "Rajasthan", "Tehsil")]),
    ("MadhyaPradesh_Villages_Pincodes.xlsx", [("23", "Madhya Pradesh", "Tehsil")]),
    ("UttarPradesh_Villages_Pincodes.xlsx", [("9", "Uttar Pradesh", "Tehsil")]),
    ("Bihar_Villages_Pincodes.xlsx", [("10", "Bihar", "Block")]),
    ("WestBengal_Villages_Pincodes.xlsx", [("19", "West Bengal", "Block")]),
    ("Punjab_Villages_Pincodes.xlsx", [("3", "Punjab", "Tehsil")]),
    ("Haryana_Villages_Pincodes.xlsx", [("6", "Haryana", "Tehsil")]),
    ("Jharkhand_Villages_Pincodes.xlsx", [("20", "Jharkhand", "Block")]),
    ("Chhattisgarh_Villages_Pincodes.xlsx", [("22", "Chhattisgarh", "Tehsil")]),
    ("Assam_Villages_Pincodes.xlsx", [("18", "Assam", "Circle")]),
    ("HimachalPradesh_Villages_Pincodes.xlsx", [("2", "Himachal Pradesh", "Tehsil")]),
    ("Uttarakhand_Villages_Pincodes.xlsx", [("5", "Uttarakhand", "Tehsil")]),
    ("JammuKashmir_Villages_Pincodes.xlsx", [("1", "J&K", "Tehsil")]),
    ("Goa_Villages_Pincodes.xlsx", [("30", "Goa", "Taluka")]),
    ("Delhi_Villages_Pincodes.xlsx", [("7", "Delhi", "Tehsil")]),
    ("Tripura_Villages_Pincodes.xlsx", [("16", "Tripura", "Subdivision")]),
    ("Manipur_Villages_Pincodes.xlsx", [("14", "Manipur", "Subdivision")]),
    ("Meghalaya_Villages_Pincodes.xlsx", [("17", "Meghalaya", "Block")]),
    ("Mizoram_Villages_Pincodes.xlsx", [("15", "Mizoram", "Block")]),
    ("ArunachalPradesh_Villages_Pincodes.xlsx", [("12", "Arunachal Pradesh", "Circle")]),
    ("Nagaland_Villages_Pincodes.xlsx", [("13", "Nagaland", "Circle")]),
]


def load(name):
    return pd.read_csv(SRC / f"{name}.{DATE}.csv", dtype=str).fillna("")


villages = load("villages")
pins = load("pincode_villages")[["Village Code", "Pincode"]].drop_duplicates("Village Code")
districts = load("districts")
subdistricts = load("subdistricts")


def village_sheet(code, sub):
    v = villages[villages["State Code"] == code].merge(pins, on="Village Code", how="left")
    df = pd.DataFrame({
        "State": v["State Name(In English)"],
        "District": v["District Name (In English)"],
        sub: v["Sub-District Name (In English)"],
        "Village": v["Village Name (In English)"],
        "Pincode": v["Pincode"].fillna(""),
        "Village (Local Name)": v["Village Name (In Local)"],
        "Village Status": v["Village Status"],
        "Village Category": v["Village Category"],
        "District LGD Code": v["District Code"],
        f"{sub} LGD Code": v["Sub-District Code"],
        "Village LGD Code": v["Village Code"],
        "Census 2011 Code": v["Census 2011 Code"],
    })
    df = df.sort_values(["District", sub, "Village"], key=lambda c: c.str.lower()).reset_index(drop=True)
    df.insert(0, "S.No.", range(1, len(df) + 1))
    return df


def subdistrict_sheet(code, sub, vdf):
    s = subdistricts[subdistricts["State Code"] == code]
    counts = vdf.groupby(f"{sub} LGD Code").size()
    df = pd.DataFrame({
        "State": s["State Name"],
        "District": s["District Name"],
        sub: s["Sub-district Name"],
        "No. of Villages": s["Sub-district Code"].map(counts).fillna(0).astype(int),
        "District LGD Code": s["District Code"],
        f"{sub} LGD Code": s["Sub-district Code"],
    })
    df = df.sort_values(["District", sub], key=lambda c: c.str.lower() if c.dtype == object else c).reset_index(drop=True)
    df.insert(0, "S.No.", range(1, len(df) + 1))
    return df


def to_numbers(df):
    for c in df.columns:
        if c.endswith("Code") or c == "Pincode":
            df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
    return df


def style(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 2000)), 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)


for fname, states in WORKBOOKS:
    if ONLY and fname not in ONLY:
        continue
    sheets, summary = {}, []
    for code, prefix, sub in states:
        vdf = village_sheet(code, sub)
        sdf = subdistrict_sheet(code, sub, vdf)
        sheets[f"{prefix} Villages"] = to_numbers(vdf)
        sheets[f"{prefix} {sub}s"] = to_numbers(sdf)
        summary.append({
            "State": vdf["State"].iloc[0],
            "Districts": int((districts["State Code"] == code).sum()),
            f"{sub}s": len(sdf),
            "Villages": len(vdf),
            "Villages with Pincode": int(vdf["Pincode"].notna().sum()),
            "Distinct Pincodes": int(vdf["Pincode"].nunique()),
        })
    summary_df = pd.DataFrame(summary)
    note = pd.DataFrame({"Source": [
        f"Local Government Directory (lgdirectory.gov.in), Govt. of India - data as on {DATE}",
        "Pincodes: LGD village-to-pincode mapping (blank = not mapped in LGD)",
    ]})
    with pd.ExcelWriter(OUT / fname, engine="openpyxl") as xw:
        summary_df.to_excel(xw, sheet_name="Summary", index=False)
        note.to_excel(xw, sheet_name="Summary", index=False, startrow=len(summary_df) + 3)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        for ws in xw.book.worksheets:
            style(ws)
    print(fname)
    print(summary_df.to_string(index=False))
