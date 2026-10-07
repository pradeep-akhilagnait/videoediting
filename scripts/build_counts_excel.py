"""State-wise counts of districts, mandals and villages (from the linked workbooks) in one Excel sheet.

Usage: python scripts/build_counts_excel.py
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

CSV = Path(__file__).resolve().parent.parent / "data" / "csv"
OUT = CSV / "State_Wise_Counts.xlsx"


def counts(path):
    x = pd.read_excel(path, sheet_name=None)
    st = next(v for k, v in x.items() if k.startswith("2 "))
    d = x["3 Districts"].rename(columns={"state_id (link)": "state_id"})
    m = next(v for k, v in x.items() if k.startswith("4 ")).rename(columns={"district_id (link)": "district_id"})
    v = x["5 Villages"].rename(columns={"mandal_id (link)": "mandal_id"})
    m = m.rename(columns={m.columns[0]: "mandal_id", m.columns[1]: "mandal_name"})
    m = m.merge(d[["district_id", "state_id"]], on="district_id")
    v = v.merge(m[["mandal_id", "state_id"]], on="mandal_id")
    r = st[["state_id", "state_name"]].set_index("state_id")
    r["Districts"] = d.groupby("state_id").size()
    r["Mandals"] = m.groupby("state_id").size()
    r["Mandal called"] = m.groupby("state_id")["mandal_type"].first() if "mandal_type" in m else "Mandal"
    r["Villages"] = v.groupby("state_id").size()
    return r


r = pd.concat([counts(CSV / "ap" / "ap_linked.xlsx"), counts(CSV / "all_states_except_ap_linked.xlsx")])
r = r.fillna({"Districts": 0, "Mandals": 0, "Villages": 0}).sort_index().reset_index()
for c in ["Districts", "Mandals", "Villages"]:
    r[c] = r[c].astype(int)
r = r.rename(columns={"state_id": "State ID", "state_name": "State / UT"})
r.insert(0, "S.No.", range(1, len(r) + 1))

ap = r["State / UT"] == "Andhra Pradesh"
totals = pd.DataFrame([
    {"S.No.": "", "State ID": "", "State / UT": "ALL INDIA (36 States/UTs)", "Mandal called": "",
     **{c: int(r[c].sum()) for c in ["Districts", "Mandals", "Villages"]}},
    {"S.No.": "", "State ID": "", "State / UT": "Andhra Pradesh only", "Mandal called": "",
     **{c: int(r.loc[ap, c].sum()) for c in ["Districts", "Mandals", "Villages"]}},
    {"S.No.": "", "State ID": "", "State / UT": "All except Andhra Pradesh (35)", "Mandal called": "",
     **{c: int(r.loc[~ap, c].sum()) for c in ["Districts", "Mandals", "Villages"]}},
])
out = pd.concat([r, pd.DataFrame([{}]), totals], ignore_index=True)
out = out[["S.No.", "State ID", "State / UT", "Districts", "Mandals", "Mandal called", "Villages"]]

with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    out.to_excel(xw, sheet_name="State Wise Counts", index=False)
    ws = xw.sheets["State Wise Counts"]
    ws.freeze_panes = "A2"
    widths = [7, 9, 44, 11, 11, 15, 12]
    for i, w in enumerate(widths, 1):
        h = ws.cell(row=1, column=i)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor="1F4E78")
        h.alignment = Alignment(horizontal="center")
        ws.column_dimensions[h.column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        if row[6].value is not None:
            row[6].number_format = "#,##0"
        for c in (row[3], row[4]):
            if c.value is not None:
                c.number_format = "#,##0"
        if row[2].value == "Andhra Pradesh":
            for c in row:
                c.fill = PatternFill("solid", fgColor="E2EFDA")
    for row in ws.iter_rows(min_row=ws.max_row - 2):
        for c in row:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FFF2CC")
    note = ws.max_row + 2
    ws.cell(row=note, column=3, value="Source: Local Government Directory (lgdirectory.gov.in), data as on 02 Oct 2026")
print(out.tail(4).to_string(index=False))
