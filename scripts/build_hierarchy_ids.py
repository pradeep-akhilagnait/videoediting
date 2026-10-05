"""Build Country > State > District > Mandal > Village tables with linked hierarchical IDs.

Each level is numbered 1, 2, 3 ... alphabetically within its parent. The full code joins the
chain with fixed widths, e.g. 1-01-01-01-0001 = India / State 1 / District 1 / Mandal 1 / Village 1,
so it is unique across the country. Every row also carries its parent's full code (the link).

Usage: python scripts/build_hierarchy_ids.py
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DATA = Path(__file__).resolve().parent.parent / "data"
COUNTRY_ID, COUNTRY = 1, "India"
# (state id, source workbook, villages sheet, sub-districts sheet, sub-district label, output)
STATES = [
    (1, "AP_Telangana_Villages_Pincodes.xlsx", "Telangana Villages", "Telangana Mandals", "Mandal",
     "Telangana_Hierarchy_IDs.xlsx"),
]


def code(*parts):
    """Fixed-width code: country 1 digit, state/district/sub-district 2 digits, village 4 digits."""
    widths = [1, 2, 2, 2, 4]
    return "-".join(str(p).zfill(w) for p, w in zip(parts, widths))


def number_within(df, parent_cols, name_col, id_col):
    """Add id_col = 1, 2, 3 ... alphabetically (case-insensitive) within each parent."""
    df = df.assign(_k=df[name_col].str.lower()).sort_values(parent_cols + ["_k"], kind="stable").drop(columns="_k")
    df[id_col] = df.groupby(parent_cols, sort=False).cumcount() + 1
    return df


def style(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 2000)), 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)


for state_id, src, vsheet, ssheet, sub, out in STATES:
    v = pd.read_excel(DATA / src, sheet_name=vsheet, dtype=str)
    s = pd.read_excel(DATA / src, sheet_name=ssheet, dtype=str)
    state = s["State"].iloc[0]
    sc = code(COUNTRY_ID, state_id)

    # Districts
    d = s[["District", "District LGD Code"]].drop_duplicates("District LGD Code")
    d = number_within(d.assign(_p=1), ["_p"], "District", "District ID").drop(columns="_p")
    d["District Code"] = [code(COUNTRY_ID, state_id, i) for i in d["District ID"]]
    dmap = d.set_index("District LGD Code")[["District ID", "District Code"]]

    # Mandals
    m = s[["District LGD Code", sub, f"{sub} LGD Code"]].copy()
    m = m.join(dmap, on="District LGD Code")
    m = number_within(m, ["District ID"], sub, f"{sub} ID")
    m[f"{sub} Code"] = [code(COUNTRY_ID, state_id, a, b) for a, b in zip(m["District ID"], m[f"{sub} ID"])]
    mmap = m.set_index(f"{sub} LGD Code")[["District ID", "District Code", f"{sub} ID", f"{sub} Code"]]

    # Villages
    vv = v.join(mmap, on=f"{sub} LGD Code")
    vv = number_within(vv, ["District ID", f"{sub} ID"], "Village", "Village ID")
    vv["Village Code"] = [code(COUNTRY_ID, state_id, a, b, c)
                          for a, b, c in zip(vv["District ID"], vv[f"{sub} ID"], vv["Village ID"])]
    vv = vv.sort_values("Village Code")

    countries = pd.DataFrame({"Country ID": [COUNTRY_ID], "Country Code": [str(COUNTRY_ID)], "Country": [COUNTRY]})
    states = pd.DataFrame({"State ID": [state_id], "State Code": [sc], "State": [state],
                           "Country Code (Parent)": [str(COUNTRY_ID)]})
    districts = pd.DataFrame({
        "District ID": d["District ID"], "District Code": d["District Code"], "District": d["District"],
        "State Code (Parent)": sc, "District LGD Code": d["District LGD Code"],
    }).sort_values("District Code")
    mandals = pd.DataFrame({
        f"{sub} ID": m[f"{sub} ID"], f"{sub} Code": m[f"{sub} Code"], sub: m[sub],
        "District Code (Parent)": m["District Code"], f"{sub} LGD Code": m[f"{sub} LGD Code"],
    }).sort_values(f"{sub} Code")
    villages = pd.DataFrame({
        "Village ID": vv["Village ID"], "Village Code": vv["Village Code"], "Village": vv["Village"],
        f"{sub} Code (Parent)": vv[f"{sub} Code"], "Pincode": vv["Pincode"],
        "Village (Local Name)": vv["Village (Local Name)"], "Village Status": vv["Village Status"],
        "Village LGD Code": vv["Village LGD Code"],
    })
    combined = pd.DataFrame({
        "Village Code": vv["Village Code"],
        "Country ID": COUNTRY_ID, "Country": COUNTRY,
        "State ID": state_id, "State": state,
        "District ID": vv["District ID"], "District": vv["District"],
        f"{sub} ID": vv[f"{sub} ID"], sub: vv[sub],
        "Village ID": vv["Village ID"], "Village": vv["Village"],
        "Pincode": vv["Pincode"],
        "State Code": sc, "District Code": vv["District Code"], f"{sub} Code": vv[f"{sub} Code"],
    })

    for df in (districts, mandals, villages, combined):
        for c in df.columns:
            if c.endswith("LGD Code") or c == "Pincode":
                df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")

    guide = pd.DataFrame({"How the IDs work": [
        f"Full code = Country-State-District-{sub}-Village, e.g. {villages['Village Code'].iloc[0]}",
        "Each ID counts 1, 2, 3 ... alphabetically inside its parent; the full code is unique across India.",
        "Widths: Country 1 digit, State 2, District 2, " + sub + " 2, Village 4.",
        "Link: every row's '(Parent)' column holds the full code of the row above it.",
        f"Example: Villages whose '{sub} Code (Parent)' = {mandals[sub + ' Code'].iloc[0]} are the villages of "
        f"{mandals[sub].iloc[0]}.",
        "Source: Local Government Directory (lgdirectory.gov.in), data as on 02Oct2026.",
    ]})
    with pd.ExcelWriter(DATA / out, engine="openpyxl") as xw:
        guide.to_excel(xw, sheet_name="Read Me", index=False)
        countries.to_excel(xw, sheet_name="1 Country", index=False)
        states.to_excel(xw, sheet_name="2 States", index=False)
        districts.to_excel(xw, sheet_name="3 Districts", index=False)
        mandals.to_excel(xw, sheet_name=f"4 {sub}s", index=False)
        villages.to_excel(xw, sheet_name="5 Villages", index=False)
        combined.to_excel(xw, sheet_name="All Levels", index=False)
        for ws in xw.book.worksheets:
            style(ws)
        xw.book["Read Me"].column_dimensions["A"].width = 110

    # sanity checks
    assert villages["Village Code"].is_unique and mandals[f"{sub} Code"].is_unique
    assert set(villages[f"{sub} Code (Parent)"]) <= set(mandals[f"{sub} Code"])
    assert set(mandals["District Code (Parent)"]) <= set(districts["District Code"])
    print(out, "| districts", len(districts), "|", sub.lower() + "s", len(mandals), "| villages", len(villages))
    print(combined.head(4).to_string(index=False))
