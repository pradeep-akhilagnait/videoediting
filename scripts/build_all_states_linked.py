"""All states/UTs except Andhra Pradesh in one linked workbook (same layout as ap_linked.xlsx).

Tabs: 1 Country, 2 States, 3 Districts, 4 Mandals, 5 Villages. Each row carries its parent's id
in a "(link)" column. Ids continue after Andhra Pradesh (already issued separately): Telangana keeps
the ids it was given, then the remaining states follow in state_id order.

Usage: python scripts/build_all_states_linked.py
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

DATA = Path(__file__).resolve().parent.parent / "data"
OUT = DATA / "csv" / "all_states_except_ap_linked.xlsx"
COUNTRY_ID, COUNTRY = 1, "India"
SUB_TYPES = ["Mandal", "Tehsil", "Taluk", "Taluka", "Block", "Circle", "Subdivision"]

# state_id as in the app's list (alphabetical, pre-2020: AP = 2, Telangana = 32). Dadra & Nagar Haveli
# and Daman & Diu were merged in 2020 and keep DNH's id 8 (9 = old Daman & Diu, unused); Ladakh is new -> 37.
STATE_IDS = {
    "Andaman And Nicobar Islands": 1, "Andhra Pradesh": 2, "Arunachal Pradesh": 3, "Assam": 4, "Bihar": 5,
    "Chandigarh": 6, "Chhattisgarh": 7, "The Dadra And Nagar Haveli And Daman And Diu": 8, "Delhi": 10,
    "Goa": 11, "Gujarat": 12, "Haryana": 13, "Himachal Pradesh": 14, "Jammu And Kashmir": 15,
    "Jharkhand": 16, "Karnataka": 17, "Kerala": 18, "Lakshadweep": 19, "Madhya Pradesh": 20,
    "Maharashtra": 21, "Manipur": 22, "Meghalaya": 23, "Mizoram": 24, "Nagaland": 25, "Odisha": 26,
    "Puducherry": 27, "Punjab": 28, "Rajasthan": 29, "Sikkim": 30, "Tamil Nadu": 31, "Telangana": 32,
    "Tripura": 33, "Uttar Pradesh": 34, "Uttarakhand": 35, "West Bengal": 36, "Ladakh": 37,
}
AP_LAST = {"district": 28, "sub": 688, "village": 17957}  # last ids issued in ap_linked.xlsx


def sort_ci(df, cols):
    return df.sort_values(cols, key=lambda c: c.str.lower() if c.dtype == object else c, kind="stable")


# LGD district list: includes districts with no sub-districts (Mumbai, Kolkata)
lgd_districts = pd.read_csv(DATA / "source" / "lgd_districts_02Oct2026.csv", dtype=str)

villages, subs = {}, {}
for f in sorted(DATA.glob("*_Villages_Pincodes.xlsx")):
    if f.name.startswith("All_India"):
        continue
    for name, df in pd.read_excel(f, sheet_name=None, dtype=str).items():
        if name == "Summary" or name.endswith("PIN Codes"):
            continue
        sub = next(t for t in SUB_TYPES if t in df.columns)
        df = df.rename(columns={sub: "Sub", f"{sub} LGD Code": "Sub LGD Code"}).assign(SubType=sub)
        state = df["State"].iloc[0]
        (villages if name.endswith(" Villages") else subs)[state] = df

states = sorted((s for s in subs if s != "Andhra Pradesh"),
                key=lambda s: (s != "Telangana", STATE_IDS[s]))
nxt = {k: v + 1 for k, v in AP_LAST.items()}
st_rows, d_rows, m_rows, v_rows = [], [], [], []
for state in states:
    sid = STATE_IDS[state]
    s, v = subs[state], villages.get(state)
    st_rows.append({"state_id": sid, "state_name": state, "country_id": COUNTRY_ID})

    ld = lgd_districts[lgd_districts["State Name (In English)"] == state]
    ld = ld.rename(columns={"District Name(In English)": "District", "District Code": "District LGD Code"})
    d = pd.concat([s[["District", "District LGD Code"]], ld[["District", "District LGD Code"]]])
    d = sort_ci(d.drop_duplicates("District LGD Code"), ["District"])
    d = d.assign(id=range(nxt["district"], nxt["district"] + len(d)))
    dmap = d.set_index("District LGD Code")["id"]
    m = s.assign(district_id=s["District LGD Code"].map(dmap))
    m = sort_ci(m, ["district_id", "Sub"])
    m = m.assign(id=range(nxt["sub"], nxt["sub"] + len(m)))
    mmap = m.set_index("Sub LGD Code")["id"]
    d_rows.append(pd.DataFrame({"district_id": d["id"], "district_name": d["District"], "state_id": sid}))
    m_rows.append(pd.DataFrame({"mandal_id": m["id"], "mandal_name": m["Sub"], "mandal_type": m["SubType"],
                                "district_id": m["district_id"]}))
    nxt.update(district=d["id"].max() + 1, sub=m["id"].max() + 1)
    if v is not None:
        vv = sort_ci(v.assign(mandal_id=v["Sub LGD Code"].map(mmap)), ["mandal_id", "Village"])
        vv = vv.assign(id=range(nxt["village"], nxt["village"] + len(vv)))
        assert vv["mandal_id"].notna().all()
        v_rows.append(pd.DataFrame({"village_id": vv["id"], "village_name": vv["Village"],
                                    "mandal_id": vv["mandal_id"], "pincode": vv["Pincode"]}))
        nxt["village"] = vv["id"].max() + 1

tabs = {
    "1 Country": pd.DataFrame({"country_id": [COUNTRY_ID], "country_name": [COUNTRY]}),
    "2 States": pd.DataFrame(st_rows).sort_values("state_id"),
    "3 Districts": pd.concat(d_rows, ignore_index=True),
    "4 Mandals": pd.concat(m_rows, ignore_index=True),
    "5 Villages": pd.concat(v_rows, ignore_index=True),
}
tabs["5 Villages"]["pincode"] = pd.to_numeric(tabs["5 Villages"]["pincode"], errors="coerce").astype("Int64")

# links must all resolve
assert set(tabs["3 Districts"]["state_id"]) <= set(tabs["2 States"]["state_id"])
assert set(tabs["4 Mandals"]["district_id"]) <= set(tabs["3 Districts"]["district_id"])
assert set(tabs["5 Villages"]["mandal_id"]) <= set(tabs["4 Mandals"]["mandal_id"])
for t, c in [("3 Districts", "district_id"), ("4 Mandals", "mandal_id"), ("5 Villages", "village_id")]:
    assert tabs[t][c].is_unique

link_fill = PatternFill("solid", fgColor="FFF2CC")
with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    for name, df in tabs.items():
        ids = [c for c in df.columns if c.endswith("_id")]
        df = df.rename(columns={c: f"{c} (link)" for c in ids[1:]})
        df.to_excel(xw, sheet_name=name, index=False)
        ws = xw.sheets[name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(df.columns, 1):
            head = ws.cell(row=1, column=i)
            head.font = Font(bold=True, color="FFFFFF")
            head.fill = PatternFill("solid", fgColor="1F4E78")
            ws.column_dimensions[head.column_letter].width = max(14, min(40, len(c) + 4,
                int(df[c].head(2000).astype(str).str.len().max()) + 4))
            if "(link)" in c:
                for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell.fill = link_fill
for k, df in tabs.items():
    idc = df.columns[0]
    print(f"{k}: {len(df)} rows, {idc} {df[idc].min()}-{df[idc].max()}")
